"""Create office-document PDF previews without changing the originals."""

import argparse
import copy
import csv
import hashlib
import json
import os
import platform
import shutil
import subprocess
import tempfile
import textwrap
import warnings
from pathlib import Path

EXTENSIONS = {".doc", ".docx", ".rtf", ".xls", ".xlsx"}
EXCLUSIONS = {
    "data/new-haven-2004/source-materials/survey-waves.xlsx": "Three sheets of respondent-level coded survey answers, not documentation."
}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def prepare_workbook(source, scratch, base_command):
    """Adjust only a temporary workbook's print layout; retain cell content."""
    try:
        import openpyxl
    except ImportError as error:
        raise SystemExit(
            "Workbook previews require: python3 -m pip install openpyxl"
        ) from error
    directory = scratch / "workbook"
    directory.mkdir()
    if source.suffix.lower() == ".xls":
        subprocess.run(
            base_command
            + ["--convert-to", "xlsx", "--outdir", str(directory), str(source)],
            check=True,
            capture_output=True,
            timeout=180,
        )
        intermediate = directory / (source.stem + ".xlsx")
        if not intermediate.is_file():
            raise SystemExit(f"LibreOffice produced no XLSX conversion: {source}")
    else:
        intermediate = directory / source.name
        shutil.copyfile(source, intermediate)
    with warnings.catch_warnings(record=True) as loading_warnings:
        warnings.simplefilter("always")
        workbook = openpyxl.load_workbook(intermediate)
    before_cells = {
        sheet.title: {
            (cell.coordinate, cell.data_type): cell.value
            for row in sheet
            for cell in row
            if cell.value is not None
        }
        for sheet in workbook
    }
    details = []
    for sheet in workbook:
        occupied = [cell for row in sheet for cell in row if cell.value is not None]
        if not occupied:
            details.append({"sheet": sheet.title, "empty": True})
            continue
        rows = max(cell.row for cell in occupied)
        columns = max(cell.column for cell in occupied)
        sheet.sheet_state = "visible"
        lengths = {}
        for column in range(1, columns + 1):
            samples = sorted(
                len(str(cell.value).split("\n")[0])
                for cell in occupied
                if cell.column == column
            )
            typical = samples[int((len(samples) - 1) * 0.8)] if samples else 10
            lengths[column] = min(45, max(12, typical + 3))
        factor = min(1, 185 / sum(lengths.values()))
        widths = {column: max(10, width * factor) for column, width in lengths.items()}
        for column, width in widths.items():
            name = openpyxl.utils.get_column_letter(column)
            sheet.column_dimensions[name].width = width
            sheet.column_dimensions[name].hidden = False
        merged_widths = {}
        for region in sheet.merged_cells.ranges:
            merged_widths[(region.min_row, region.min_col)] = sum(
                widths.get(column, 12)
                for column in range(region.min_col, region.max_col + 1)
            )
        heights = {row: 16 for row in range(1, rows + 1)}
        for cell in occupied:
            alignment = copy.copy(cell.alignment)
            alignment.wrap_text = True
            alignment.vertical = "top"
            alignment.shrink_to_fit = False
            cell.alignment = alignment
            font = copy.copy(cell.font)
            font.sz = min(font.sz or 10, 10)
            cell.font = font
            width = merged_widths.get((cell.row, cell.column), widths[cell.column])
            lines = sum(
                max(1, len(textwrap.wrap(line, max(1, int(width - 3)))))
                for line in str(cell.value).split("\n")
            )
            heights[cell.row] = max(heights[cell.row], lines * 13 + 5)
        for row, height in heights.items():
            sheet.row_dimensions[row].height = height
            sheet.row_dimensions[row].hidden = False
        sheet.print_area = f"A1:{openpyxl.utils.get_column_letter(columns)}{rows}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = (
            sheet.PAPERSIZE_A3 if max(heights.values()) <= 700 else "66"
        )
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins.left = sheet.page_margins.right = 0.25
        sheet.page_margins.top = sheet.page_margins.bottom = 0.4
        sheet.row_breaks = openpyxl.worksheet.pagebreak.RowBreak()
        sheet.col_breaks = openpyxl.worksheet.pagebreak.ColBreak()
        details.append(
            {
                "sheet": sheet.title,
                "rows": rows,
                "columns": columns,
                "formulas": sum(cell.data_type == "f" for cell in occupied),
                "images": len(sheet._images),
                "charts": len(sheet._charts),
                "rows_over_700_points": [r for r, h in heights.items() if h > 700],
                "maximum_row_height": max(heights.values()),
                "conversion_warnings": [str(w.message) for w in loading_warnings],
            }
        )
    workbook.save(intermediate)
    checked = openpyxl.load_workbook(intermediate)
    after_cells = {
        sheet.title: {
            (cell.coordinate, cell.data_type): cell.value
            for row in sheet
            for cell in row
            if cell.value is not None
        }
        for sheet in checked
    }
    if before_cells != after_cells:
        raise ValueError(f"Temporary layout changed workbook content: {source}")
    return intermediate, details


def main():
    parser = argparse.ArgumentParser(
        description=__doc__
        + " Requires LibreOffice, Poppler; workbooks also require openpyxl."
    )
    parser.add_argument("--report", required=True, type=Path)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    executable = shutil.which("soffice")
    if not executable:
        raise SystemExit("LibreOffice is required: install it and add soffice to PATH.")
    for command in ("pdfinfo", "pdftotext"):
        if not shutil.which(command):
            raise SystemExit(f"Poppler is required: {command} is not on PATH.")
    fontconfig = Path("/opt/homebrew/etc/fonts/fonts.conf")
    if fontconfig.exists():
        os.environ.setdefault("FONTCONFIG_FILE", str(fontconfig))
    version = subprocess.check_output([executable, "--version"], text=True).strip()
    ledger_path = root / "metadata/document_previews.csv"
    if ledger_path.exists():
        with ledger_path.open(newline="") as stream:
            ledger = {row["source_path"]: row for row in csv.DictReader(stream)}
    else:
        ledger = {}
    report = {
        "converter": version,
        "platform": platform.platform(),
        "excluded": EXCLUSIONS,
        "previews": [],
    }
    with tempfile.TemporaryDirectory(prefix="dp-document-previews-") as temporary:
        scratch = Path(temporary)
        profile = (scratch / "profile").as_uri()
        for source in sorted((root / "data").rglob("*")):
            relative = source.relative_to(root).as_posix()
            if source.suffix.lower() not in EXTENSIONS or relative in EXCLUSIONS:
                continue
            checksum = sha256(source)
            registered = ledger.get(relative)
            if registered:
                preview = root / registered["preview_path"]
                if (
                    registered["source_sha256"] != checksum
                    or not preview.exists()
                    or registered["preview_sha256"] != sha256(preview)
                ):
                    raise SystemExit(
                        f"Registered preview changed; review first: {relative}"
                    )
                report["previews"].append(dict(registered, status="unchanged"))
                continue
            destination = source.with_suffix(".pdf")
            if destination.exists():
                destination = source.with_name(source.name + ".pdf")
            if destination.exists():
                raise SystemExit(f"Unregistered preview already exists: {destination}")
            source_scratch = Path(tempfile.mkdtemp(prefix="source-", dir=scratch))
            output = source_scratch / "output"
            output.mkdir()
            base_command = [
                executable,
                f"-env:UserInstallation={profile}",
                "--headless",
            ]
            conversion_source = source
            workbook_details = []
            if source.suffix.lower() in {".xls", ".xlsx"}:
                conversion_source, workbook_details = prepare_workbook(
                    source, source_scratch, base_command
                )
            command = base_command + [
                "--convert-to",
                "pdf",
                "--outdir",
                str(output),
                str(conversion_source),
            ]
            subprocess.run(command, check=True, capture_output=True, timeout=180)
            converted = output / (source.stem + ".pdf")
            if not converted.is_file():
                raise SystemExit(f"LibreOffice produced no PDF conversion: {relative}")
            info = subprocess.check_output(["pdfinfo", str(converted)], text=True)
            text = subprocess.check_output(
                ["pdftotext", str(converted), "-"], text=True
            )
            pages = next(
                int(line.split(":", 1)[1])
                for line in info.splitlines()
                if line.startswith("Pages:")
            )
            if pages < 1 or not text.strip():
                raise SystemExit(f"Empty PDF conversion: {relative}")
            if sha256(source) != checksum:
                raise SystemExit(f"Original changed during conversion: {relative}")
            shutil.copyfile(converted, destination)
            report["previews"].append(
                {
                    "source_path": relative,
                    "preview_path": destination.relative_to(root).as_posix(),
                    "source_sha256": checksum,
                    "preview_sha256": sha256(destination),
                    "converter_version": version,
                    "pages": pages,
                    "font_configuration": os.environ.get(
                        "FONTCONFIG_FILE", "system-default"
                    ),
                    "text_characters": len(text),
                    "command": command,
                    "status": "converted-needs-visual-review",
                    "workbook_layout": workbook_details,
                }
            )
            args.report.write_text(json.dumps(report, indent=2) + "\n")
    args.report.write_text(json.dumps(report, indent=2) + "\n")


if __name__ == "__main__":
    main()
